import sys
import types
sys.modules['imghdr'] = types.ModuleType('imghdr')
import logging
import os
from openpyxl import load_workbook
from telegram import Update
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    filters,
    ContextTypes,
)
 
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
EXCEL_FILE = "students.xlsx"
WAITING_ID = 1
 
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)
 
 
def load_data():
    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
 
    ws_s = wb["students"]
    rows_s = list(ws_s.iter_rows(values_only=True))
    headers_s = [str(h).strip() if h else "" for h in rows_s[0]]
    students = []
    for row in rows_s[1:]:
        record = {headers_s[i]: (str(row[i]).strip() if row[i] is not None else "")
                  for i in range(len(headers_s))}
        students.append(record)
 
    ws_e = wb["exam"]
    rows_e = list(ws_e.iter_rows(values_only=True))
    headers_e = [str(h).strip() if h else "" for h in rows_e[1]]
    exams = []
    for row in rows_e[2:]:
        record = {headers_e[i]: (str(row[i]).strip() if row[i] is not None else "")
                  for i in range(len(headers_e))}
        exams.append(record)
 
    wb.close()
    return students, exams
 
 
try:
    students_data, exams_data = load_data()
    logger.info(f"Data loaded: {len(students_data)} students | {len(exams_data)} exams")
except Exception as e:
    logger.error(f"Excel load error: {e}")
    raise
 
 
def get_schedule(student_id: str) -> str:
    sid = student_id.strip()
    student_courses = [r for r in students_data if r.get("رقم المتدرب", "") == sid]
 
    if not student_courses:
        return "لم يتم العثور على رقم تدريبي مطابق.\nتأكد من الرقم وأعد المحاولة."
 
    student_name = student_courses[0].get("اسم المتدرب", "")
    lines = [
        f"اسم المتدرب: {student_name}",
        f"الرقم التدريبي: {sid}",
        "━━━━━━━━━━━━━━━━━━━━",
        "جدول الاختبارات النهائي:",
        "",
    ]
 
    for course in student_courses:
        ref_num     = course.get("الرقم المرجعي", "").strip()
        course_name = course.get("المقرر", "")
        exam_type   = course.get("نوع الاختبار", "")
        stu_status  = course.get("حالة المتدرب", "")
        crs_status  = course.get("حالة المقرر", "")
 
        exam_rows = [e for e in exams_data if e.get("الرقم المرجعي", "").strip() == ref_num]
 
        if not exam_rows:
            lines += [
                f"المقرر: {course_name}",
                f"الرقم المرجعي: {ref_num}",
                f"نوع الاختبار: {exam_type}",
                f"حالة المتدرب: {stu_status}",
                f"حالة المقرر: {crs_status}",
                "لا يوجد موعد في جدول الاختبارات",
                "",
            ]
        else:
            for exam in exam_rows:
                date      = exam.get("التاريخ", "")
                day       = exam.get("اليوم", "")
                period    = exam.get("الفترة", "")
                location  = exam.get("موقع اللجنة", "")
                committee = exam.get("اللجنة", "")
 
                if location.startswith("=") or not location:
                    location = "غير محدد"
 
                lines += [
                    f"المقرر: {course_name}",
                    f"الرقم المرجعي: {ref_num}",
                    f"التاريخ: {date} | {day}",
                    f"الفترة: {period}",
                    f"موقع اللجنة: {location}",
                    f"رقم اللجنة: {committee}",
                    f"حالة المتدرب: {stu_status}",
                    f"حالة المقرر: {crs_status}",
                    "",
                ]
 
    lines += ["━━━━━━━━━━━━━━━━━━━━", "نظام جداول الاختبارات"]
    return "\n".join(lines)
 
 
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "اهلا بك في نظام جداول الاختبارات\n\n"
        "اكتب: جدول\n"
        "للاستعلام عن جدول اختباراتك النهائي"
    )
 
 
async def request_schedule(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("من فضلك ادخل رقمك التدريبي:")
    return WAITING_ID
 
 
async def receive_id(update: Update, context: ContextTypes.DEFAULT_TYPE):
    student_id = update.message.text.strip()
 
    if not student_id.isdigit():
        await update.message.reply_text("الرقم التدريبي يجب ان يكون ارقاما فقط. اعد الادخال:")
        return WAITING_ID
 
    await update.message.reply_text("جاري البحث...")
    result = get_schedule(student_id)
    await update.message.reply_text(result)
    return ConversationHandler.END
 
 
async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("تم الغاء الطلب.")
    return ConversationHandler.END
 
 
async def fallback_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    keywords = ["جدول", "اختبار", "جدولي", "مواعيد", "اختباراتي", "schedule"]
    if any(k in text for k in keywords):
        await update.message.reply_text("من فضلك ادخل رقمك التدريبي:")
        return WAITING_ID
    await update.message.reply_text("اكتب: جدول\nللاستعلام عن جدول اختباراتك.")
    return ConversationHandler.END
 
 
def main():
    if not BOT_TOKEN:
        raise ValueError("BOT_TOKEN environment variable is not set!")
 
    app = ApplicationBuilder().token(BOT_TOKEN).build()
 
    conv_handler = ConversationHandler(
        entry_points=[
            CommandHandler("jadwal", request_schedule),
            MessageHandler(filters.TEXT & ~filters.COMMAND, fallback_handler),
        ],
        states={
            WAITING_ID: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_id)
            ],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
    )
 
    app.add_handler(CommandHandler("start", start))
    app.add_handler(conv_handler)
 
    logger.info("Bot is running...")
    app.run_polling()
 
 
if __name__ == "__main__":
    main()
from flask import Flask
import threading

app_web = Flask(__name__)

@app_web.route("/")
def home():
    return "Bot is running ✅"

def run_web():
    port = int(os.environ.get("PORT", 10000))
    app_web.run(host="0.0.0.0", port=port)

# تشغيل السيرفر مع البوت
if __name__ == "__main__":
    threading.Thread(target=run_web).start()
    main()
